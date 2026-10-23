"""
Text from a document — THE one way LakeLogic turns a file into text.

``extract_text(path)`` picks the reader by file type and returns what it read, which reader read
it, and — when there is nothing to read — says so instead of returning an empty string that looks
like success. Pipelines and the LakeLogic platform (e.g. documents attached to an assistant) both
call this, so a file that works in one works in the other.

Readers (all pip-only, imported only when that type is read):

=========================  ==================================  ====================================
Type                       Reader                              Install
=========================  ==================================  ====================================
.txt .md .csv .json .yaml  built-in (UTF-8)                    —
.html .htm                 built-in ``html.parser``            —
.eml                       built-in ``email``                  —
.pdf                       ``pdfplumber`` (text + tables)      ``lakelogic[documents]``
scanned .pdf               ``pdfplumber`` renders, RapidOCR    ``lakelogic[documents]``
.png .jpg .jpeg .webp .bmp RapidOCR (ONNX, no PyTorch)         ``lakelogic[documents]``
.docx                      ``python-docx`` (text + tables)     ``lakelogic[documents]``
.pptx                      ``python-pptx``                     ``lakelogic[documents]``
.xlsx                      ``openpyxl``                        ``lakelogic[documents]``
.doc .ppt .xls             refused: re-save as .docx/.pptx/.xlsx (needs LibreOffice to read)
=========================  ==================================  ====================================

OCR reads printed words, not meaning: a chart or diagram gives little text, and the result says
``no_text`` rather than inventing any.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Union

PathLike = Union[str, Path]

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".yaml", ".yml", ".log"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
LEGACY_OFFICE = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | IMAGE_SUFFIXES | {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm", ".eml"}

INSTALL_HINT = "pip install 'lakelogic[documents]'"


class ExtractionUnavailable(ImportError):
    """The reader for this type is not installed; the message says what to install."""


class UnsupportedDocument(ValueError):
    """This type cannot be read; the message says why and what to do instead."""


@dataclass
class ExtractedText:
    text: str
    #: Which reader produced it (``plain``, ``pdfplumber``, ``rapidocr``, ``python-docx``, ...).
    extractor: str
    #: Tables found, each a list of rows (lists of cell strings). PDF, Word and Excel only.
    tables: List[List[List[str]]] = field(default_factory=list)
    #: Pages / slides / sheets read, when the format has them.
    units: Optional[int] = None
    #: True when the file was read but held no text (e.g. a photo of a chart). Not an error.
    no_text: bool = False
    #: One plain sentence about anything notable (OCR used, pages without text, ...).
    note: Optional[str] = None

    def as_plain_text(self) -> str:
        """The text, with each table appended as tab-separated rows — what a reader of the
        document (person or model) should see."""
        parts = [self.text] if self.text else []
        for i, table in enumerate(self.tables, 1):
            parts.append(f"[Table {i}]\n" + "\n".join("\t".join(c or "" for c in row) for row in table))
        return "\n\n".join(parts)


def extract_text(source: PathLike, *, filename: Optional[str] = None) -> ExtractedText:
    """Read a document and return its text. ``source`` is a path; ``filename`` overrides the name
    used to pick the reader (e.g. an upload saved under a temporary name)."""
    path = Path(source)
    suffix = Path(filename or path.name).suffix.lower()
    if suffix in LEGACY_OFFICE:
        raise UnsupportedDocument(
            f"{suffix} is an older Office format that needs LibreOffice to read. "
            f"Save it as {LEGACY_OFFICE[suffix]} and try again."
        )
    if suffix in TEXT_SUFFIXES:
        return _plain(path)
    if suffix in {".html", ".htm"}:
        return _html(path)
    if suffix == ".eml":
        return _email(path)
    if suffix == ".pdf":
        return _pdf(path)
    if suffix in IMAGE_SUFFIXES:
        return _image(path)
    if suffix == ".docx":
        return _docx(path)
    if suffix == ".pptx":
        return _pptx(path)
    if suffix == ".xlsx":
        return _xlsx(path)
    raise UnsupportedDocument(f"{suffix or 'This file type'} is not a document LakeLogic reads.")


def _done(text: str, extractor: str, **kw) -> ExtractedText:
    text = (text or "").strip()
    tables = kw.pop("tables", None) or []
    return ExtractedText(text=text, extractor=extractor, tables=tables, no_text=not text and not tables, **kw)


# ── plain text ──────────────────────────────────────────────────────────────
def _plain(path: Path) -> ExtractedText:
    data = path.read_bytes()
    if b"\x00" in data:
        raise UnsupportedDocument(f"{path.name} is not a text file.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise UnsupportedDocument(f"{path.name} is not UTF-8 text.") from None
    return _done(text, "plain")


def _html(path: Path) -> ExtractedText:
    from html.parser import HTMLParser

    class _Text(HTMLParser):
        skip = {"script", "style", "head"}

        def __init__(self):
            super().__init__()
            self.out: List[str] = []
            self._depth = 0

        def handle_starttag(self, tag, attrs):
            if tag in self.skip:
                self._depth += 1

        def handle_endtag(self, tag):
            if tag in self.skip and self._depth:
                self._depth -= 1

        def handle_data(self, data):
            if not self._depth and data.strip():
                self.out.append(data.strip())

    p = _Text()
    p.feed(path.read_text(encoding="utf-8", errors="replace"))
    return _done("\n".join(p.out), "html")


def _email(path: Path) -> ExtractedText:
    from email import policy
    from email.parser import BytesParser

    msg = BytesParser(policy=policy.default).parse(io.BytesIO(path.read_bytes()))
    head = "\n".join(f"{k}: {msg[k]}" for k in ("From", "To", "Date", "Subject") if msg[k])
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body is not None else ""
    if body is not None and body.get_content_type() == "text/html":
        tmp = path.with_suffix(".tmp_lakelogic.html")
        try:
            tmp.write_text(text, encoding="utf-8")
            text = _html(tmp).text
        finally:
            tmp.unlink(missing_ok=True)
    return _done(f"{head}\n\n{text}", "email")


# ── PDF (text first; OCR only for pages with none) ──────────────────────────
def _pdf(path: Path) -> ExtractedText:
    try:
        import pdfplumber
    except ImportError as e:
        raise ExtractionUnavailable(f"Reading PDF needs pdfplumber: {INSTALL_HINT}") from e
    texts: List[str] = []
    tables: List[List[List[str]]] = []
    blank_pages: List[int] = []
    with pdfplumber.open(str(path)) as doc:
        pages = len(doc.pages)
        for i, page in enumerate(doc.pages, 1):
            t = (page.extract_text() or "").strip()
            if t:
                texts.append(t)
            else:
                blank_pages.append(i)
            for table in page.extract_tables() or []:
                tables.append([[("" if c is None else str(c)) for c in row] for row in table])
        note = None
        if blank_pages:
            # Scanned pages: render and OCR them, if OCR is installed. Otherwise say so.
            try:
                ocr = _rapidocr()
            except ExtractionUnavailable:
                note = (
                    f"{len(blank_pages)} of {pages} page(s) have no text layer (scanned); "
                    f"install OCR to read them: {INSTALL_HINT}"
                )
            else:
                read = 0
                for i in blank_pages:
                    img = doc.pages[i - 1].to_image(resolution=200).original
                    t = _ocr_image(ocr, img)
                    if t:
                        texts.append(t)
                        read += 1
                note = f"{len(blank_pages)} scanned page(s) read with OCR ({read} had text)."
    return _done("\n\n".join(texts), "pdfplumber", tables=tables, units=pages, note=note)


# ── images (OCR) ────────────────────────────────────────────────────────────
def _rapidocr():
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError as e:
        raise ExtractionUnavailable(f"Reading text in images needs RapidOCR: {INSTALL_HINT}") from e
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


_OCR_ENGINE = None


def _ocr_image(engine, image) -> str:
    """OCR one image (a path, bytes, or a PIL image). Lines in reading order."""
    try:
        import numpy as np

        if hasattr(image, "convert"):  # PIL image → array
            image = np.array(image.convert("RGB"))
    except ImportError:
        pass
    result, _ = engine(image)
    return "\n".join(line[1] for line in (result or []) if len(line) > 1 and line[1]).strip()


def _image(path: Path) -> ExtractedText:
    text = _ocr_image(_rapidocr(), str(path))
    res = _done(text, "rapidocr")
    if res.no_text:
        res.note = "No text found in this image (OCR reads printed words, not charts or pictures)."
    return res


# ── Office ──────────────────────────────────────────────────────────────────
def _docx(path: Path) -> ExtractedText:
    try:
        import docx
    except ImportError as e:
        raise ExtractionUnavailable(f"Reading Word needs python-docx: {INSTALL_HINT}") from e
    d = docx.Document(str(path))
    text = "\n".join(p.text for p in d.paragraphs if p.text.strip())
    tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in d.tables]
    return _done(text, "python-docx", tables=tables)


def _pptx(path: Path) -> ExtractedText:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise ExtractionUnavailable(f"Reading PowerPoint needs python-pptx: {INSTALL_HINT}") from e
    deck = Presentation(str(path))
    slides = []
    for n, slide in enumerate(deck.slides, 1):
        lines = [
            s.text_frame.text.strip()
            for s in slide.shapes
            if getattr(s, "has_text_frame", False) and s.text_frame.text.strip()
        ]
        if lines:
            slides.append(f"[Slide {n}]\n" + "\n".join(lines))
    return _done("\n\n".join(slides), "python-pptx", units=len(deck.slides))


def _xlsx(path: Path) -> ExtractedText:
    try:
        import openpyxl
    except ImportError as e:
        raise ExtractionUnavailable(f"Reading Excel needs openpyxl: {INSTALL_HINT}") from e
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    tables, names = [], []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if rows:
            tables.append(rows)
            names.append(ws.title)
    wb.close()
    return _done("Sheets: " + ", ".join(names) if names else "", "openpyxl", tables=tables, units=len(names))


def tables_as_csv(result: ExtractedText) -> List[str]:
    """Each table as CSV text — for writing tables out as their own files."""
    out = []
    for table in result.tables:
        buf = io.StringIO()
        csv.writer(buf).writerows(table)
        out.append(buf.getvalue())
    return out
