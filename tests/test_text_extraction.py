"""lakelogic.extract_text — one reader per document type, each proven on a real generated file.

Each test builds the document with the library a user's tool would (reportlab, Pillow,
python-docx, python-pptx, openpyxl), reads it back with extract_text, and checks the CONTENT,
not just that something came back.
"""

from __future__ import annotations

import pytest

from lakelogic import extract_text
from lakelogic.core.text_extraction import UnsupportedDocument

CODE = "PELICAN 42"


def _font(size=40):
    from PIL import ImageFont

    for name in ("arial.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _text_image(path, lines):
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (1000, 120 * len(lines) + 80), "white")
    d = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        d.text((40, 40 + i * 120), line, fill="black", font=_font())
    img.save(path)
    return img


# ── plain text and friends ──────────────────────────────────────────────────
def test_plain_text_and_markdown(tmp_path):
    f = tmp_path / "note.md"
    f.write_text(f"# Note\n\nThe code word is {CODE}.\n", encoding="utf-8")
    r = extract_text(f)
    assert r.extractor == "plain" and CODE in r.text and not r.no_text


def test_binary_pretending_to_be_text_is_refused(tmp_path):
    f = tmp_path / "x.txt"
    f.write_bytes(b"\x00\x01\x02")
    with pytest.raises(UnsupportedDocument, match="not a text file"):
        extract_text(f)


def test_html_drops_scripts_and_styles(tmp_path):
    f = tmp_path / "p.html"
    f.write_text(
        f"<html><head><style>x{{}}</style></head><body><h1>Title</h1><script>evil()</script>"
        f"<p>Code {CODE}</p></body></html>",
        encoding="utf-8",
    )
    r = extract_text(f)
    assert CODE in r.text and "evil" not in r.text and "Title" in r.text


def test_email_keeps_headers_and_body(tmp_path):
    f = tmp_path / "m.eml"
    f.write_bytes((f"From: a@x.com\nTo: b@x.com\nSubject: Hello\n\nThe code is {CODE}.\n").encode())
    r = extract_text(f)
    assert "Subject: Hello" in r.text and CODE in r.text


# ── PDF ─────────────────────────────────────────────────────────────────────
def test_pdf_text_and_table(tmp_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    f = tmp_path / "report.pdf"
    doc = SimpleDocTemplate(str(f), pagesize=A4)
    table = Table([["rider_id", "city"], ["R1", "London"], ["R2", "Paris"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc.build([Paragraph(f"Quarterly note. Code {CODE}.", getSampleStyleSheet()["Normal"]), table])

    r = extract_text(f)
    assert r.extractor == "pdfplumber" and r.units == 1
    assert CODE in r.text
    assert ["rider_id", "city"] in r.tables[0] and ["R2", "Paris"] in r.tables[0]
    assert "R2\tParis" in r.as_plain_text()


def test_scanned_pdf_is_read_with_ocr(tmp_path):
    img = _text_image(tmp_path / "scan.png", [f"Invoice code {CODE}"])
    f = tmp_path / "scan.pdf"
    img.save(f, "PDF", resolution=150)  # an image only: no text layer
    r = extract_text(f)
    assert "PELICAN" in r.text.upper()
    assert r.note and "OCR" in r.note


# ── images ──────────────────────────────────────────────────────────────────
def test_image_text_with_rapidocr(tmp_path):
    f = tmp_path / "shot.png"
    _text_image(f, [f"Error code {CODE}", "Table silver_trips failed"])
    r = extract_text(f)
    assert r.extractor == "rapidocr"
    assert "PELICAN" in r.text.upper() and "silver" in r.text.lower()


def test_an_image_without_text_says_so(tmp_path):
    from PIL import Image

    f = tmp_path / "blank.png"
    Image.new("RGB", (400, 300), "white").save(f)
    r = extract_text(f)
    assert r.no_text and "No text found" in (r.note or "")


# ── Office ──────────────────────────────────────────────────────────────────
def test_word_text_and_table(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Design note", 1)
    d.add_paragraph(f"The code word is {CODE}.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "field", "type"
    t.cell(1, 0).text, t.cell(1, 1).text = "trip_id", "string"
    f = tmp_path / "note.docx"
    d.save(f)
    r = extract_text(f)
    assert r.extractor == "python-docx" and CODE in r.text
    assert r.tables == [[["field", "type"], ["trip_id", "string"]]]


def test_powerpoint_slides(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    for title in ("Roadmap", f"Code {CODE}"):
        s = deck.slides.add_slide(deck.slide_layouts[1])
        s.shapes.title.text = title
    f = tmp_path / "deck.pptx"
    deck.save(f)
    r = extract_text(f)
    assert r.extractor == "python-pptx" and r.units == 2
    assert "[Slide 2]" in r.text and CODE in r.text


def test_excel_sheets_as_tables(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trips"
    ws.append(["trip_id", "fare"])
    ws.append(["T1", 12.5])
    wb.create_sheet("Empty")
    f = tmp_path / "book.xlsx"
    wb.save(f)
    r = extract_text(f)
    assert r.extractor == "openpyxl" and r.units == 1  # the empty sheet is not a table
    assert r.tables == [[["trip_id", "fare"], ["T1", "12.5"]]]
    assert "Trips" in r.text


# ── refusals ────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("name,modern", [("old.doc", ".docx"), ("old.ppt", ".pptx"), ("old.xls", ".xlsx")])
def test_legacy_office_is_refused_with_what_to_do(tmp_path, name, modern):
    f = tmp_path / name
    f.write_bytes(b"x")
    with pytest.raises(UnsupportedDocument, match=modern):
        extract_text(f)


def test_unknown_type_is_refused(tmp_path):
    f = tmp_path / "a.xyz"
    f.write_bytes(b"x")
    with pytest.raises(UnsupportedDocument):
        extract_text(f)


def test_filename_overrides_the_saved_name(tmp_path):
    f = tmp_path / "upload.tmp"
    f.write_text(f"code {CODE}", encoding="utf-8")
    assert CODE in extract_text(f, filename="note.txt").text
