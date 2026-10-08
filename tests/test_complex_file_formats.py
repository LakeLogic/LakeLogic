"""Nested JSON/XML and awkward Excel through a contract (found 2026-10-07).

- flatten_nested used to turn an array into `items_values_values_values_values_values`, so
  `explode` found nothing; arrays now stay whole, under their own name, as JSON text.
- XML nested elements arrive in the same shape as nested JSON.
- Excel honours source.options sheet_name / header_row / skip_footer.
"""

import json

import openpyxl
import pytest

from lakelogic import DataProcessor

pl = pytest.importorskip("polars")

ITEMS_CONTRACT = {
    "transformations": [
        {"explode": {"field": "items", "output": "item"}},
        {"json_extract": {"field": "sku", "source": "item", "path": "$.sku"}},
        {"json_extract": {"field": "qty", "source": "item", "path": "$.qty", "cast": "int"}},
    ],
    "model": {
        "fields": [
            {"name": "customer_address_city", "type": "string", "required": True},
            {"name": "sku", "type": "string", "required": True},
            {"name": "qty", "type": "int", "required": True},
        ]
    },
    "quality": {"row_rules": [{"name": "qty_positive", "sql": "qty > 0"}]},
}


def _run(path, fmt, **source):
    contract = {
        "version": "1.0",
        "info": {"title": "t"},
        "source": {"type": "landing", "path": str(path), "format": fmt, **source},
        **ITEMS_CONTRACT,
    }
    return DataProcessor(engine="polars", contract=contract).run_source()


def test_nested_json_arrays_explode_into_one_row_per_item(tmp_path):
    p = tmp_path / "o.json"
    p.write_text(
        json.dumps(
            [
                {
                    "id": 1,
                    "customer": {"address": {"city": "London"}},
                    "items": [{"sku": "A", "qty": 2}, {"sku": "B", "qty": 1}],
                },
                {"id": 2, "customer": {"address": {"city": "Rome"}}, "items": [{"sku": "C", "qty": -1}]},
            ]
        )
    )
    good, bad = _run(p, "json", flatten_nested=True)
    assert sorted(good["sku"].to_list()) == ["A", "B"]
    assert good["customer_address_city"].to_list() == ["London", "London"]
    assert not any("_values" in c for c in good.columns)
    assert bad["sku"].to_list() == ["C"]


def test_nested_xml_reads_like_nested_json(tmp_path):
    p = tmp_path / "o.xml"
    p.write_text(
        '<export xmlns="urn:x"><orders>'
        '<order id="1"><customer><address><city>London</city></address></customer>'
        "<items><item><sku>A</sku><qty>2</qty></item><item><sku>B</sku><qty>1</qty></item></items></order>"
        '<order id="2"><customer><address><city>Rome</city></address></customer>'
        "<items><item><sku>C</sku><qty>two</qty></item></items></order>"
        "</orders></export>"
    )
    good, bad = _run(p, "xml", flatten_nested=True)
    assert sorted(good["sku"].to_list()) == ["A", "B"]
    assert good["id"].to_list() == ["1", "1"]
    assert bad["sku"].to_list() == ["C"]  # "two" is not an int


def test_excel_reads_named_sheet_header_row_and_skips_footer(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    wb.active.append(["not data"])
    ws = wb.create_sheet("Orders")
    for row in (["Title row"], ["order_id", "amount"], [1, 10.5], [], [2, -3], ["TOTAL", 7.5]):
        ws.append(row)
    p = tmp_path / "o.xlsx"
    wb.save(p)
    contract = {
        "version": "1.0",
        "info": {"title": "x"},
        "source": {
            "type": "landing",
            "path": str(p),
            "format": "xlsx",
            "options": {"sheet_name": "Orders", "header_row": 2, "skip_footer": 1},
        },
        "model": {
            "fields": [{"name": "order_id", "type": "long", "required": True}, {"name": "amount", "type": "double"}]
        },
        "quality": {"row_rules": [{"name": "amount_not_negative", "sql": "amount >= 0"}]},
    }
    good, bad = DataProcessor(engine="polars", contract=contract).run_source()
    assert good["order_id"].to_list() == [1]
    assert bad["order_id"].to_list() == [2]  # the TOTAL row is gone, not quarantined


def test_excel_unknown_sheet_is_a_clear_error(tmp_path):
    from lakelogic.core.processor import _read_excel_polars

    wb = openpyxl.Workbook()
    p = tmp_path / "o.xlsx"
    wb.save(p)
    with pytest.raises(ValueError, match="sheet 'Nope' not found"):
        _read_excel_polars(str(p), {"sheet_name": "Nope"})


FW_LAYOUT = [
    {"name": "id", "start": 1, "width": 4},
    {"name": "name", "start": 5, "width": 6},
    {"name": "amount", "start": 11, "width": 5},
]


def _fw_contract(path, **opts):
    return {
        "version": "1.0",
        "info": {"title": "fw"},
        "source": {
            "type": "landing",
            "path": str(path),
            "format": "fixed_width",
            "options": {"columns": FW_LAYOUT, **opts},
        },
        "model": {
            "fields": [
                {"name": "id", "type": "long", "required": True},
                {"name": "name", "type": "string", "required": True},
                {"name": "amount", "type": "long"},
            ]
        },
    }


def test_fixed_width_slices_trims_and_skips_header_and_trailer(tmp_path):
    p = tmp_path / "c.txt"
    p.write_text("HDR\n0001Ada   00105\n\n0002Bo    0X200\n0003      00003\nTRL 3\n")
    good, bad = DataProcessor(engine="polars", contract=_fw_contract(p, skip_rows=1, skip_footer=1)).run_source()
    assert good.select(["id", "name", "amount"]).rows() == [(1, "Ada", 105)]
    assert sorted(bad["id"].to_list()) == [2, 3]  # bad amount; missing name


def test_fixed_width_without_header_skip_quarantines_the_header(tmp_path):
    p = tmp_path / "c.txt"
    p.write_text("HDR line here\n0001Ada   00105\n")
    good, bad = DataProcessor(engine="polars", contract=_fw_contract(p)).run_source()
    assert len(good) == 1 and len(bad) == 1


def test_fixed_width_needs_a_layout(tmp_path):
    from lakelogic.core.processor import _read_fixed_width

    p = tmp_path / "c.txt"
    p.write_text("0001Ada\n")
    with pytest.raises(ValueError, match="needs a layout"):
        _read_fixed_width(str(p), {})
    with pytest.raises(ValueError, match="start and width must be >= 1"):
        _read_fixed_width(str(p), {"columns": [{"name": "a", "start": 0, "width": 2}]})
    # A short line gives null for the fields past its end, not an error.
    df = _read_fixed_width(
        str(p), {"columns": [{"name": "id", "start": 1, "width": 4}, {"name": "tail", "start": 20, "width": 3}]}
    )
    assert df.rows() == [("0001", None)]


# ── Fixed-width, layout on the fields (range) + record_length (2026-10-07) ─────

BACS_FIELDS = [
    {"name": "record_type", "range": [0, 1], "type": "string"},
    {"name": "sort_code", "range": [1, 7], "type": "string", "required": True},
    {"name": "amount_pence", "range": [7, 15], "type": "long", "required": True},
]


def _bacs(path, **options):
    contract = {
        "version": "1.0",
        "info": {"title": "bacs"},
        "source": {"type": "landing", "path": str(path), "format": "fixed_width", "options": options},
        "model": {"fields": BACS_FIELDS},
    }
    return DataProcessor(engine="polars", contract=contract).run_source()


def test_field_ranges_and_record_length_quarantine_a_short_record(tmp_path):
    p = tmp_path / "s.dat"
    p.write_text("D20157500001050\nD30963400002000\nD404784000\n", encoding="ascii")
    good, bad = _bacs(p, record_length=15, encoding="ascii")
    assert good.select(["sort_code", "amount_pence"]).rows() == [("201575", 1050), ("309634", 2000)]
    assert bad["_lakelogic_errors"].to_list() == [["Line length mismatch: expected 15, got 10"]]
    assert not [c for c in list(good.columns) + list(bad.columns) if c.startswith("__")]


def test_a_file_with_no_line_breaks_is_split_every_record_length(tmp_path):
    p = tmp_path / "s.dat"
    p.write_bytes(b"D20157500001050D30963400002000D4047")  # mainframe style: no newlines, last one cut
    good, bad = _bacs(p, record_length=15)
    assert good["sort_code"].to_list() == ["201575", "309634"]
    assert bad["_lakelogic_errors"].to_list()[0][0] == "Line length mismatch: expected 15, got 5"


def test_ebcdic_decodes(tmp_path):
    p = tmp_path / "s.dat"
    p.write_bytes("D20157500001050".encode("cp037"))
    good, _ = _bacs(p, record_length=15, encoding="cp037")
    assert good.select(["sort_code", "amount_pence"]).rows() == [("201575", 1050)]


def test_fixed_width_settings_on_source_are_refused_with_the_fix(tmp_path):
    """OLC 0.21: format settings live only under source.options (one place for every format)."""
    contract = {
        "version": "1.0",
        "info": {"title": "bacs"},
        "source": {"type": "landing", "path": str(tmp_path), "format": "fixed_width", "record_length": 15},
        "model": {"fields": BACS_FIELDS},
    }
    with pytest.raises(Exception, match="`record_length` belongs under source.options"):
        DataProcessor(engine="polars", contract=contract)


def test_fixed_width_reads_a_cloud_object_through_fsspec(tmp_path, monkeypatch):
    fsspec = pytest.importorskip("fsspec")
    with fsspec.open("memory://landing/bacs/s.dat", "wb") as fh:
        fh.write(b"D20157500001050\n")
    from lakelogic.core import processor as proc_mod

    proc = DataProcessor(
        engine="polars",
        contract={
            "version": "1.0",
            "info": {"title": "bacs"},
            "source": {"type": "landing", "path": "memory://landing/bacs/s.dat", "format": "fixed_width"},
            "model": {"fields": BACS_FIELDS},
        },
    )
    monkeypatch.setattr(proc, "_is_uri_path", lambda p: str(p).startswith("memory://"))
    monkeypatch.setattr(proc, "_get_cloud_storage_options", lambda p: {})
    df = proc._read_fixed_width_source("memory://landing/bacs/s.dat")
    assert df.rows() == [("D", "201575", "00001050")]
    assert proc_mod.RECORD_ERROR_COLUMN not in df.columns  # no record_length, no check


@pytest.mark.parametrize(
    "rng, msg",
    [([3, 3], "must have 0 <= start < end"), ([-1, 2], "must have 0 <= start < end"), ("x", "range must be")],
)
def test_bad_ranges_are_clear_errors(tmp_path, rng, msg):
    from lakelogic.core.processor import _read_fixed_width

    p = tmp_path / "s.dat"
    p.write_text("D201575\n")
    with pytest.raises(ValueError, match=msg):
        _read_fixed_width(str(p), {}, fields=[{"name": "a", "range": rng}])
